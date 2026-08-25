"""Assemble imported real-canary inputs into the v1 runtime layout.

This helper is deliberately independent of the external canary repositories:
it reads only the inputs committed under ``tests/v2/contract/fixtures``. Phase K
must reuse it when running the v1↔v2 canary comparison so both implementations
receive the same README-derived ``data/`` tree.

Usage from the repository root::

    from pathlib import Path
    from tests.v2.compat.canary_assembly import assemble_canary_case

    data_root = assemble_canary_case("invoice", Path("/tmp/canary-invoice"))

The caller owns the destination and any cleanup. Existing ``data/`` trees are
never overwritten.
"""

from __future__ import annotations

import io
import json
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import load_workbook

CANARY_MODES = (
    "excelinvoice",
    "invoice",
    "multidatatile",
    "rdeformat",
    "smarttable",
)
CANARY_INPUT_ROOT = Path(__file__).parents[1] / "contract" / "fixtures" / "inputs" / "canary"
_WORKSHEET_XML = "xl/worksheets/sheet1.xml"
_XML_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_SYNTHETIC_USER_NAMES = tuple(f"RDE,User{index:02d}" for index in range(1, 5))
SYNTHETIC_OWNER_IDS = tuple(f"{index:056d}" for index in range(1, 5))
SYNTHETIC_SAMPLE_OWNER_ID = f"{5:056d}"
_CANARY_OWNER_PATHS = (
    Path("excelinvoice/invoice/invoice_auto.json"),
    Path("multidatatile/invoice/invoice.json"),
    Path("smarttable/invoice/invoice.json"),
)


def sanitize_canary_inputs(inputs_root: Path = CANARY_INPUT_ROOT) -> None:
    """Deterministically remove personal identities from an imported canary copy.

    This is the final step of the real-canary import pipeline. User-list row
    order, rather than the source names or hashes, defines four stable
    replacements: ``RDE,User01`` through ``RDE,User04`` and the 56-digit
    synthetic IDs ending in ``1`` through ``4``. Registration cells retain the
    first user's identity and cached lookup result. The three canary invoice
    data owners use the first synthetic ID, while their sample owner uses the
    non-colliding synthetic ID ending in ``5``. Reapplying the sanitizer is
    byte-stable, so a future source refresh cannot silently reintroduce
    personal data.

    Workbook member payloads are rewritten inside the existing OOXML archive.
    This preserves formulas and their cached values, unlike an openpyxl save.
    Personal author and absolute-path metadata are normalized at the same
    boundary even though the runtime does not consume them. Workbook and JSON
    transformations are fully rendered and validated before any file is
    written, so a late failure leaves the complete imported corpus unchanged.

    Args:
        inputs_root: Imported canary root containing all five mode families.

    Raises:
        FileNotFoundError: If a required workbook or invoice is absent.
        ValueError: If identities are malformed or cross-file owner IDs differ.
    """
    workbook_path = inputs_root / "excelinvoice/inputdata/cb550_excelinvoice.xlsx"
    required = (workbook_path, *(inputs_root / path for path in _CANARY_OWNER_PATHS))
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)

    source_names, source_ids = _read_excelinvoice_identities(workbook_path)
    invoice_owner_ids = tuple(
        _read_invoice_identity(inputs_root / relative, "basic", "dataOwnerId")
        for relative in _CANARY_OWNER_PATHS
    )
    if len(set(invoice_owner_ids)) != 1:
        msg = "unexpected dataOwnerId mismatch across imported canary invoices"
        raise ValueError(msg)
    sample_owner_ids = tuple(
        _read_invoice_identity(inputs_root / relative, "sample", "ownerId")
        for relative in _CANARY_OWNER_PATHS
    )
    if len(set(sample_owner_ids)) != 1:
        msg = "unexpected sample.ownerId mismatch across imported canary invoices"
        raise ValueError(msg)

    rewritten: dict[Path, bytes] = {
        workbook_path: _render_sanitized_excelinvoice(
            workbook_path,
            source_names,
            source_ids,
        ),
    }
    for relative in _CANARY_OWNER_PATHS:
        path = inputs_root / relative
        payload = path.read_text(encoding="utf-8")
        payload = _replace_invoice_identity(
            path,
            "dataOwnerId",
            invoice_owner_ids[0],
            SYNTHETIC_OWNER_IDS[0],
            payload,
        )
        payload = _replace_invoice_identity(
            path,
            "ownerId",
            sample_owner_ids[0],
            SYNTHETIC_SAMPLE_OWNER_ID,
            payload,
        )
        rewritten[path] = payload.encode()

    for path, payload in rewritten.items():
        if path.read_bytes() != payload:
            path.write_bytes(payload)


def _read_excelinvoice_identities(
    workbook: Path | bytes,
) -> tuple[tuple[str, ...], tuple[str, ...]]:
    """Validate and return the workbook's four user-list identities."""
    cached_source = io.BytesIO(workbook) if isinstance(workbook, bytes) else workbook
    formula_source = io.BytesIO(workbook) if isinstance(workbook, bytes) else workbook
    cached_book = load_workbook(cached_source, read_only=True, data_only=True)
    formula_book = load_workbook(formula_source, read_only=True, data_only=False)
    try:
        users = cached_book["ユーザーリスト"]
        registration = cached_book["登録用"]
        formulas = formula_book["登録用"]
        names = tuple(users[f"A{row}"].value for row in range(2, 6))
        owner_ids = tuple(users[f"B{row}"].value for row in range(2, 6))
        if not all(isinstance(name, str) and name for name in names) or len(set(names)) != 4:
            msg = "ExcelInvoice user-list names must be four distinct text values"
            raise ValueError(msg)
        if not all(
            isinstance(owner_id, str) and len(owner_id) == 56 and owner_id.isalnum()
            for owner_id in owner_ids
        ) or len(set(owner_ids)) != 4:
            msg = "ExcelInvoice user-list IDs must be four distinct 56-character values"
            raise ValueError(msg)
        if any(registration[cell].value != names[0] for cell in ("C5", "C6", "J5", "J6")):
            msg = "ExcelInvoice registration names must reference user-list row 2"
            raise ValueError(msg)
        if any(registration[cell].value != owner_ids[0] for cell in ("D5", "D6", "K5", "K6")):
            msg = "ExcelInvoice registration ID caches must reference user-list row 2"
            raise ValueError(msg)
        if any(not str(formulas[cell].value).startswith("=IFERROR(VLOOKUP(") for cell in ("D5", "D6", "K5", "K6")):
            msg = "ExcelInvoice registration ID formulas are missing"
            raise ValueError(msg)
        return names, owner_ids
    finally:
        cached_book.close()
        formula_book.close()


def _read_invoice_identity(path: Path, section: str, key: str) -> str:
    """Return one validated identity from an imported invoice."""
    invoice = json.loads(path.read_text(encoding="utf-8"))
    owner_id = invoice.get(section, {}).get(key)
    if not isinstance(owner_id, str) or len(owner_id) != 56 or not owner_id.isalnum():
        msg = f"unexpected {section}.{key} in {path.as_posix()}"
        raise ValueError(msg)
    return owner_id


def _replace_invoice_identity(
    path: Path,
    key: str,
    source: str,
    replacement: str,
    payload: str,
) -> str:
    """Replace one key-scoped invoice identity without otherwise reformatting JSON.

    The replacement targets the exact ``"key": "value"`` pair so the same
    56-character identity may legitimately appear under another key (for
    example one person acting as both ``dataOwnerId`` and ``sample.ownerId``)
    without breaking the single-occurrence guard or touching the wrong field.
    """
    if source == replacement:
        return payload
    pattern = re.compile(
        rf'("{re.escape(key)}"\s*:\s*)"{re.escape(source)}"',
    )
    if len(pattern.findall(payload)) != 1:
        msg = f"unexpected {key} occurrence count in {path.as_posix()}"
        raise ValueError(msg)
    return pattern.sub(
        lambda match: f'{match.group(1)}"{replacement}"',
        payload,
    )


def _render_sanitized_excelinvoice(
    workbook_path: Path,
    source_names: tuple[str, ...],
    source_ids: tuple[str, ...],
) -> bytes:
    """Render sanitized OOXML bytes while retaining formula caches."""
    replacements = {
        **dict(zip(source_names, _SYNTHETIC_USER_NAMES, strict=True)),
        **dict(zip(source_ids, SYNTHETIC_OWNER_IDS, strict=True)),
    }
    effective = {
        source.encode(): replacement.encode()
        for source, replacement in replacements.items()
        if source != replacement
    }
    output = io.BytesIO()
    changed = False
    with zipfile.ZipFile(workbook_path) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            payload = source.read(member.filename)
            rewritten = payload
            for old, new in effective.items():
                rewritten = rewritten.replace(old, new)
            if member.filename == "xl/workbook.xml":
                rewritten = re.sub(
                    rb'url="/Users/[^\"]*/[^\"]*/"',
                    b'url="/Users/rde-user/Downloads/"',
                    rewritten,
                )
            elif member.filename == "docProps/core.xml":
                rewritten = _sanitize_core_properties(rewritten)
            changed = changed or rewritten != payload
            target.writestr(member, rewritten)
    rendered = output.getvalue() if changed else workbook_path.read_bytes()
    _assert_sanitized_excelinvoice(rendered)
    return rendered


def _sanitize_core_properties(payload: bytes) -> bytes:
    """Replace personal OOXML author metadata with a stable marker."""
    root = ET.fromstring(payload)  # noqa: S314
    for element in root.iter():
        if element.tag.rsplit("}", maxsplit=1)[-1] in {"creator", "lastModifiedBy"}:
            element.text = "RDE"
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _assert_sanitized_excelinvoice(workbook: Path | bytes) -> None:
    """Verify sanitized workbook values and cached formulas after OOXML rewrite."""
    names, owner_ids = _read_excelinvoice_identities(workbook)
    if names != _SYNTHETIC_USER_NAMES or owner_ids != SYNTHETIC_OWNER_IDS:
        msg = "ExcelInvoice identity sanitization did not produce canonical values"
        raise ValueError(msg)


def assemble_canary_case(
    mode: str,
    destination: Path,
    *,
    inputs_root: Path = CANARY_INPUT_ROOT,
) -> Path:
    """Materialize one imported canary family below ``destination/data``.

    The imported directories preserve each canary README's source naming. The
    ExcelInvoice ``invoice_auto.json`` is renamed at assembly time. Its workbook
    is also normalized to ``*_excel_invoice.xlsx`` because v1 detects this mode
    only through that filename suffix; the README/source name predates the
    convention and is abbreviated. The material's second registration row has
    a leading-space typo in its two filename cells, while the matching
    ``input.zip`` member has no space. Assembly trims only those cells so v1 can
    register the row; already-corrected upstream material is left untouched.
    README mode 3 uses ``tasksupport2`` as shorthand, but the canary material
    shows that it contains only the differing MultiDataTile ``rdeconfig.yaml``;
    assembly therefore copies the base schema/metadata files and overlays that
    config. These corrections record the material's demonstrated intent rather
    than silently selecting another mode or omitting required schemas. An empty
    ``unpacked`` directory is recreated because Git cannot retain it.

    Args:
        mode: One of the five values in :data:`CANARY_MODES`.
        destination: Empty case root that will receive ``data/``.
        inputs_root: Repository-owned canary input root, overridable by tests.

    Returns:
        The materialized ``data`` directory.

    Raises:
        ValueError: If ``mode`` is not a supported canary mode.
        FileExistsError: If ``destination/data`` already exists.
    """
    if mode not in CANARY_MODES:
        supported = ", ".join(CANARY_MODES)
        msg = f"unsupported canary mode {mode!r}; expected one of: {supported}"
        raise ValueError(msg)

    source_root = inputs_root / mode
    data_root = destination / "data"
    if mode == "multidatatile":
        data_root.mkdir(parents=True)
        shutil.copytree(source_root / "inputdata", data_root / "inputdata")
        shutil.copytree(source_root / "invoice", data_root / "invoice")
        shutil.copytree(source_root / "tasksupport", data_root / "tasksupport")
        shutil.copy2(
            source_root / "tasksupport2" / "rdeconfig.yaml",
            data_root / "tasksupport" / "rdeconfig.yaml",
        )
    else:
        shutil.copytree(source_root, data_root)
    if mode == "excelinvoice":
        (data_root / "invoice" / "invoice_auto.json").rename(
            data_root / "invoice" / "invoice.json",
        )
        workbook_path = data_root / "inputdata" / "cb550_excel_invoice.xlsx"
        (data_root / "inputdata" / "cb550_excelinvoice.xlsx").rename(workbook_path)
        _trim_excelinvoice_filename_cells(workbook_path)
    (data_root / "unpacked").mkdir()
    return data_root


def _trim_excelinvoice_filename_cells(workbook_path: Path) -> None:
    """Trim the two mistyped row-2 filename cells in the assembled copy."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["登録用"]
    values = {coordinate: sheet[coordinate].value for coordinate in ("A6", "U6")}
    workbook.close()
    if not all(isinstance(value, str) for value in values.values()):
        msg = "ExcelInvoice filename cells A6 and U6 must contain text"
        raise ValueError(msg)
    trimmed = {coordinate: value.strip() for coordinate, value in values.items()}
    if values == trimmed:
        return
    if len(set(trimmed.values())) != 1:
        msg = "ExcelInvoice filename cells A6 and U6 must identify the same raw file"
        raise ValueError(msg)
    _rewrite_excelinvoice_filename_cells(workbook_path, trimmed)


def _rewrite_excelinvoice_filename_cells(
    workbook_path: Path,
    values: dict[str, str],
) -> None:
    """Rewrite two worksheet values without invalidating formula caches."""
    output = io.BytesIO()
    with zipfile.ZipFile(workbook_path) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            payload = source.read(member.filename)
            if member.filename == _WORKSHEET_XML:
                payload = _rewrite_registration_sheet(payload, values)
            target.writestr(member, payload)
    workbook_path.write_bytes(output.getvalue())


def _rewrite_registration_sheet(payload: bytes, values: dict[str, str]) -> bytes:
    """Return registration-sheet XML with only A6 and U6 values replaced."""
    namespace = f"{{{_XML_NAMESPACE}}}"
    ET.register_namespace("", _XML_NAMESPACE)
    # The payload comes only from the repository-owned canary workbook copy.
    root = ET.fromstring(payload)  # noqa: S314
    cells = {
        cell.attrib["r"]: cell
        for cell in root.iter(f"{namespace}c")
        if cell.attrib.get("r") in values
    }
    if cells.keys() != values.keys():
        msg = "ExcelInvoice registration cells A6 and U6 were not found"
        raise ValueError(msg)

    cached_value = cells["A6"].find(f"{namespace}v")
    if cached_value is None:
        msg = "ExcelInvoice registration cell A6 has no cached formula value"
        raise ValueError(msg)
    cached_value.text = values["A6"]

    data_name = cells["U6"]
    for child in list(data_name):
        if child.tag in {f"{namespace}v", f"{namespace}is"}:
            data_name.remove(child)
    data_name.set("t", "inlineStr")
    inline_string = ET.SubElement(data_name, f"{namespace}is")
    ET.SubElement(inline_string, f"{namespace}t").text = values["U6"]
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
