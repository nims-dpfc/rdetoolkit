"""Tests for the reusable real-canary assembly helper.

Equivalence partitions prepared before implementation:

| Input | Partition | Expected | Test ID |
| --- | --- | --- | --- |
| canary mode | each of five supported modes | README-compatible ``data/`` tree | TC-H4-ASSEMBLY-001 |
| ExcelInvoice source | ``invoice_auto.json`` | materialized as ``invoice.json`` | TC-H4-ASSEMBLY-002 |
| ExcelInvoice source | abbreviated workbook name | materialized with ``_excel_invoice`` suffix | TC-H4-ASSEMBLY-006 |
| ExcelInvoice source | row-2 filename typo | only the two filename cells are whitespace-trimmed | TC-H4-ASSEMBLY-007 |
| imported canary copy | real user names and IDs | deterministic synthetic values replace PII | TC-H4-ASSEMBLY-008 |
| imported canary copy | already sanitized | second sanitization is byte-identical | TC-H4-ASSEMBLY-009 |
| imported canary copy | required invoice missing | ``FileNotFoundError`` before partial rewrite | TC-H4-ASSEMBLY-010 |
| imported canary copy | unexpected owner ID | ``ValueError`` before partial rewrite | TC-H4-ASSEMBLY-011 |
| canary mode | unsupported name | ``ValueError`` | TC-H4-ASSEMBLY-003 |

Boundary values prepared before implementation:

| Input | Boundary | Expected | Test ID |
| --- | --- | --- | --- |
| destination | empty directory | assembly succeeds | TC-H4-ASSEMBLY-001 |
| destination | existing ``data/`` | ``FileExistsError`` | TC-H4-ASSEMBLY-004 |
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tests.v2.compat.canary_assembly import (
    CANARY_INPUT_ROOT,
    CANARY_MODES,
    assemble_canary_case,
    sanitize_canary_inputs,
)

_SYNTHETIC_IDS = tuple(f"{index:056d}" for index in range(1, 5))


@pytest.mark.parametrize(
    ("mode", "expected_input"),
    [
        ("excelinvoice", "cb550_excel_invoice.xlsx"),
        ("invoice", "1.jpg"),
        ("multidatatile", "LiTFSI_DME_charge_discharge_dummy.csv"),
        ("rdeformat", "afm_rdeformat.zip"),
        ("smarttable", "smarttable_battery_electrolyte.csv"),
    ],
)
def test_assemble_each_canary_mode__tc_h4_assembly_001(
    tmp_path: Path,
    mode: str,
    expected_input: str,
) -> None:
    """TC-H4-ASSEMBLY-001: every imported mode becomes a runnable data tree."""
    # Given: an empty destination and one supported imported canary mode
    destination = tmp_path / mode
    # When: assembling the mode from repository-owned inputs
    data_root = assemble_canary_case(mode, destination)
    # Then: the common layout and representative real input exist
    assert data_root == destination / "data"
    assert (data_root / "inputdata" / expected_input).is_file()
    assert (data_root / "invoice" / "invoice.json").is_file()
    assert (data_root / "tasksupport" / "rdeconfig.yaml").is_file()
    assert (data_root / "unpacked").is_dir()


def test_excelinvoice_auto_invoice_is_renamed__tc_h4_assembly_002(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-002: README's invoice_auto rename is deterministic."""
    # Given: the ExcelInvoice canary source family
    # When: assembling its data tree
    data_root = assemble_canary_case("excelinvoice", tmp_path)
    # Then: only the runtime name is present
    assert (data_root / "invoice" / "invoice.json").is_file()
    assert not (data_root / "invoice" / "invoice_auto.json").exists()


def test_excelinvoice_workbook_uses_v1_detection_suffix(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-006: ExcelInvoice uses v1's required filename suffix."""
    # Given: the canary workbook whose source name predates v1's detection convention
    # When: assembling its runtime data tree
    data_root = assemble_canary_case("excelinvoice", tmp_path)
    # Then: the workbook stem ends in the exact suffix recognized by v1
    input_names = sorted(path.name for path in (data_root / "inputdata").glob("*.xlsx"))
    assert input_names == ["cb550_excel_invoice.xlsx"]


def test_excelinvoice_workbook_filename_cells_are_trimmed(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-007: assembly trims the two mistyped filename cells."""
    # Given: the canary workbook whose second registration row has leading spaces
    # When: assembling and loading the runtime workbook
    source_workbook = (
        CANARY_INPUT_ROOT / "excelinvoice/inputdata/cb550_excelinvoice.xlsx"
    )
    source_bytes = source_workbook.read_bytes()
    data_root = assemble_canary_case("excelinvoice", tmp_path)
    workbook = load_workbook(
        data_root / "inputdata" / "cb550_excel_invoice.xlsx",
        read_only=True,
        data_only=True,
    )
    sheet = workbook["登録用"]
    # Then: the raw-file name and dataName cells contain the exact ZIP member name
    assert sheet["A6"].value == "CrossBeam550_手動_1.tif"
    assert sheet["U6"].value == "CrossBeam550_手動_1.tif"
    assert source_workbook.read_bytes() == source_bytes


def test_imported_canary_copy_is_deterministically_sanitized__tc_h4_assembly_008(
    tmp_path: Path,
) -> None:
    """TC-H4-ASSEMBLY-008: the import pipeline replaces names and owner IDs."""
    # Given: a repository-owned copy of the imported canary families
    imported = tmp_path / "canary"
    shutil.copytree(CANARY_INPUT_ROOT, imported)
    # When: applying the deterministic import sanitizer
    sanitize_canary_inputs(imported)
    formula_book = load_workbook(
        imported / "excelinvoice/inputdata/cb550_excelinvoice.xlsx",
        read_only=True,
        data_only=False,
    )
    cached_book = load_workbook(
        imported / "excelinvoice/inputdata/cb550_excelinvoice.xlsx",
        read_only=True,
        data_only=True,
    )
    users = cached_book["ユーザーリスト"]
    registration = cached_book["登録用"]
    formulas = formula_book["登録用"]
    # Then: row order defines synthetic identities and registration caches retain integrity
    assert [users[f"A{row}"].value for row in range(2, 6)] == [
        "RDE,User01",
        "RDE,User02",
        "RDE,User03",
        "RDE,User04",
    ]
    assert [users[f"B{row}"].value for row in range(2, 6)] == list(_SYNTHETIC_IDS)
    assert [registration[cell].value for cell in ("C5", "C6", "J5", "J6")] == [
        "RDE,User01",
    ] * 4
    assert [registration[cell].value for cell in ("D5", "D6", "K5", "K6")] == [
        _SYNTHETIC_IDS[0],
    ] * 4
    assert formulas["D5"].value == '=IFERROR(VLOOKUP(C5,ユーザーリスト!$A:$B,2,0),"")'
    formula_book.close()
    cached_book.close()
    for relative in (
        "excelinvoice/invoice/invoice_auto.json",
        "multidatatile/invoice/invoice.json",
        "smarttable/invoice/invoice.json",
    ):
        invoice = json.loads((imported / relative).read_text(encoding="utf-8"))
        assert invoice["basic"]["dataOwnerId"] == _SYNTHETIC_IDS[0]
        assert invoice["sample"]["ownerId"] == "0" * 55 + "5"


def test_canary_sanitization_is_idempotent__tc_h4_assembly_009(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-009: re-sanitizing imported material changes no bytes."""
    # Given: one imported tree that has already passed through sanitization
    imported = tmp_path / "canary"
    shutil.copytree(CANARY_INPUT_ROOT, imported)
    sanitize_canary_inputs(imported)
    first = {
        path.relative_to(imported): path.read_bytes()
        for path in imported.rglob("*")
        if path.is_file()
    }
    # When: applying the same sanitizer after a simulated re-import
    sanitize_canary_inputs(imported)
    second = {
        path.relative_to(imported): path.read_bytes()
        for path in imported.rglob("*")
        if path.is_file()
    }
    # Then: the complete imported tree is byte-identical
    assert second == first


def test_canary_sanitization_rejects_missing_invoice__tc_h4_assembly_010(
    tmp_path: Path,
) -> None:
    """TC-H4-ASSEMBLY-010: incomplete imports fail before any rewrite."""
    # Given: an imported tree missing one required owner-bearing invoice
    imported = tmp_path / "canary"
    shutil.copytree(CANARY_INPUT_ROOT, imported)
    missing = imported / "smarttable/invoice/invoice.json"
    missing.unlink()
    workbook = imported / "excelinvoice/inputdata/cb550_excelinvoice.xlsx"
    original_workbook = workbook.read_bytes()
    # When: sanitizing the incomplete tree
    with pytest.raises(FileNotFoundError, match="smarttable/invoice/invoice.json"):
        sanitize_canary_inputs(imported)
    # Then: preflight validation prevents a partial workbook rewrite
    assert workbook.read_bytes() == original_workbook


def test_canary_sanitization_rejects_unexpected_owner_id__tc_h4_assembly_011(
    tmp_path: Path,
) -> None:
    """TC-H4-ASSEMBLY-011: unknown owner IDs cannot pass the import boundary."""
    # Given: complete material with an owner ID outside the approved mapping
    imported = tmp_path / "canary"
    shutil.copytree(CANARY_INPUT_ROOT, imported)
    invoice_path = imported / "smarttable/invoice/invoice.json"
    invoice = json.loads(invoice_path.read_text(encoding="utf-8"))
    invoice["basic"]["dataOwnerId"] = "f" * 56
    invoice_path.write_text(json.dumps(invoice), encoding="utf-8")
    workbook = imported / "excelinvoice/inputdata/cb550_excelinvoice.xlsx"
    original_workbook = workbook.read_bytes()
    # When: sanitizing an import that cannot preserve approved identity mapping
    with pytest.raises(ValueError, match="unexpected dataOwnerId"):
        sanitize_canary_inputs(imported)
    # Then: preflight validation prevents a partial workbook rewrite
    assert workbook.read_bytes() == original_workbook


def test_multidatatile_overlays_config_on_base_tasksupport(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-005: mode 3 overlays its config on the complete base support."""
    # Given: the abbreviated README mode-3 task-support source
    # When: assembling the MultiDataTile canary
    data_root = assemble_canary_case("multidatatile", tmp_path)
    # Then: required schemas coexist with the mode-selecting overlay
    assert (data_root / "tasksupport" / "invoice.schema.json").is_file()
    assert (data_root / "tasksupport" / "metadata-def.json").is_file()
    assert 'extended_mode: "MultiDataTile"' in (
        data_root / "tasksupport" / "rdeconfig.yaml"
    ).read_text(encoding="utf-8")


def test_unsupported_canary_mode_is_rejected__tc_h4_assembly_003(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-003: unsupported mode names fail before copying."""
    # Given: a mode outside the sorted supported inventory
    # When: requesting assembly
    with pytest.raises(ValueError, match="unsupported canary mode"):
        assemble_canary_case("unknown", tmp_path)
    # Then: no data tree is created
    assert not (tmp_path / "data").exists()


def test_existing_data_tree_is_not_overwritten__tc_h4_assembly_004(tmp_path: Path) -> None:
    """TC-H4-ASSEMBLY-004: assembly never overwrites an existing data tree."""
    # Given: a destination whose data directory already exists
    (tmp_path / "data").mkdir()
    # When: assembling a supported mode
    with pytest.raises(FileExistsError):
        assemble_canary_case(CANARY_MODES[0], tmp_path)
    # Then: the pre-existing directory remains intact
    assert list((tmp_path / "data").iterdir()) == []
